"""新機能の通し（ 統合 ） テスト。
``run_fetch`` → ``run_tables`` まで通して、 ``SOQL調査.xlsx`` と
``列名の対応.csv`` を検証する。 偽物の Salesforce クライアントで実 API には
繋がない。

シート構成は 3 枚 （ 一覧 / 列の対応 / 関係図 ） 。 テストの参照シート名も合わせる。
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

DOMAIN = "https://example--sandbox.sandbox.my.salesforce.com/lightning/r/Report"


# ── 偽物の Salesforce クライアント ─────────────────────────────────


class _ReportStub:
    def __init__(
        self, describe: dict, *, get_table=None, get_error: BaseException | None = None
    ) -> None:
        self._describe = describe
        self._get_table = get_table if get_table is not None else []
        self._get_error = get_error
        self.get_calls: list[str] = []

    def describe(self, report_id: str) -> dict:
        return self._describe

    def get(self, report_id: str, filters: object = None, allow_truncated: bool = False) -> object:
        self.get_calls.append(report_id)
        if self._get_error is not None:
            raise self._get_error
        return self._get_table


class _FakeClient:
    def __init__(
        self,
        describe: dict,
        object_describes: dict[str, dict] | None = None,
        *,
        get_table=None,
        get_error: BaseException | None = None,
    ) -> None:
        self._object_describes = object_describes or {}
        self.report = _ReportStub(describe, get_table=get_table, get_error=get_error)
        self.describe_object_calls: list[str] = []

    def data_path(self, path: str) -> str:
        return path

    def request(
        self,
        method: str,
        path: str,
        body: object = None,
        component: str = "other",
        headers: dict[str, str] | None = None,
        data: str | None = None,
    ) -> tuple[object, dict[str, str]]:
        return None, {}

    def describe_object(self, name: str) -> dict:
        self.describe_object_calls.append(name)
        return self._object_describes.get(name, {"name": name, "fields": []})

    def __enter__(self) -> _FakeClient:
        return self

    def __exit__(self, *args: object) -> None:
        return None


class _FakeSite:
    def __init__(self, client: _FakeClient) -> None:
        self._client = client

    def __call__(self) -> _FakeSite:
        return self

    def __enter__(self) -> _FakeClient:
        return self._client

    def __exit__(self, *args: object) -> None:
        return None


def _site_for(client: _FakeClient):
    def resolver(url: str) -> _FakeSite:
        return _FakeSite(client)

    return resolver


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        columns = next(reader)
        rows = [dict(zip(columns, line, strict=True)) for line in reader]
    return columns, rows


def _read_xlsx(path: Path) -> dict[str, list[dict[str, str]]]:
    """xlsx を ``{シート名: [行 dict, ...]}`` で読む。

    一覧 / 列の対応 シートは 1 行目=案内、 2 行目=空、 3 行目=テーブル見出し、 4 行目以降=データ
    という構造。 関係図シートは 1 行目=テーブル見出し、 2 行目以降=データ 。
    それぞれ自動的に見つける
    """
    workbook = load_workbook(path)
    result: dict[str, list[dict[str, str]]] = {}
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        all_rows = list(sheet.iter_rows(values_only=True))
        header_idx = -1
        for i, row in enumerate(all_rows):
            if not row:
                continue
            if all(isinstance(v, str) and v for v in row):
                header_idx = i
                break
        if header_idx < 0:
            result[sheet_name] = []
            continue
        headers = list(all_rows[header_idx])
        data_rows = all_rows[header_idx + 1 :]
        result[sheet_name] = [
            {
                header: ("" if value is None else value)
                for header, value in zip(headers, row, strict=False)
            }
            for row in data_rows
        ]
    workbook.close()
    return result


def _describe(
    report_type: str,
    columns: list[str],
    *,
    format_: str = "TABULAR",
    column_labels: dict[str, str] | None = None,
) -> dict:
    label_map = column_labels or {}
    return {
        "reportMetadata": {
            "reportType": {"type": report_type},
            "reportFormat": format_,
            "detailColumns": columns,
        },
        "reportExtendedMetadata": {
            "detailColumnInfo": {col: {"label": label_map.get(col, col)} for col in columns},
        },
    }


def _object_describes() -> dict[str, dict]:
    return {
        "Opportunity": {
            "name": "Opportunity",
            "fields": [
                {"name": "Name", "label": "商談名", "type": "string"},
                {"name": "Amount", "label": "金額", "type": "currency"},
                {
                    "name": "Contact__c",
                    "label": "取引先責任者",
                    "type": "reference",
                    "referenceTo": ["Contact"],
                    "relationshipName": "Contact__r",
                },
            ],
        },
        "Contact": {
            "name": "Contact",
            "fields": [
                {"name": "Name", "label": "氏名", "type": "string"},
                {"name": "Email", "label": "メール", "type": "email"},
                {"name": "Phone", "label": "電話", "type": "phone"},
            ],
        },
    }


class _Row:
    def __init__(self, count: int) -> None:
        self._count = count

    def __len__(self) -> int:
        return self._count


# ── 統合テスト ─────────────────────────────────────────────────────


def test_e2e_run_tables_writes_overview_and_columns_sheet(tmp_path: Path) -> None:
    """``run_fetch`` → ``run_tables`` で 一覧シートに 2000件超 と 個人情報 が出る。

    似たレポート判定は 1 グループになるが、 その出力先は一覧シートの 「 似たレポート 」
    列にまとめられた（ 別シートは無くなった ）。
    """
    from src import pii
    from src.fetch import RELATED_DEPTH, RELATED_MAX, ROW_LIMIT_ENABLED, run_fetch
    from src.master import MasterEntry
    from src.tables import run_tables

    describe_csv_path = tmp_path / "describe.csv"
    pii_keywords = pii.DEFAULT_KEYWORDS
    pii_person_objects = pii.DEFAULT_PERSON_OBJECTS

    from comken.exceptions import SalesforceReportTruncatedError

    client = _FakeClient(
        describe=_describe(
            "Opportunity",
            ["Opp.Name", "Contact.Email"],
            column_labels={"Contact.Email": "メール"},
        ),
        object_describes=_object_describes(),
        get_error=SalesforceReportTruncatedError("00O000000000001", 2000),
    )
    site_for = _site_for(client)

    entry_1001 = MasterEntry(
        key="1001",
        summary="テスト1001",
        url=f"{DOMAIN}/00O000000000001/view",
        enabled=True,
    )
    outcome_1001 = run_fetch(
        [entry_1001],
        related_depth=RELATED_DEPTH,
        related_max=RELATED_MAX,
        row_limit=ROW_LIMIT_ENABLED,
        pii_keywords=pii_keywords,
        pii_person_objects=pii_person_objects,
        describe_csv_path=describe_csv_path,
        site_for=site_for,
    )[0]
    assert outcome_1001.status == "ok"

    client2 = _FakeClient(
        describe=_describe(
            "Opportunity",
            ["Opp.Name", "Contact.Email"],
            column_labels={"Contact.Email": "メール"},
        ),
        object_describes=_object_describes(),
        get_table=_Row(10),
    )
    entry_1002 = MasterEntry(
        key="1002",
        summary="テスト1002",
        url=f"{DOMAIN}/00O000000000002/view",
        enabled=True,
    )
    outcome_1002 = run_fetch(
        [entry_1002],
        related_depth=RELATED_DEPTH,
        related_max=RELATED_MAX,
        row_limit=ROW_LIMIT_ENABLED,
        pii_keywords=pii_keywords,
        pii_person_objects=pii_person_objects,
        describe_csv_path=describe_csv_path,
        site_for=_site_for(client2),
    )[0]
    assert outcome_1002.status == "ok"

    output = tmp_path / "output"
    output.mkdir(parents=True, exist_ok=True)
    records = [outcome_1001.record, outcome_1002.record]
    run_tables(
        output,
        records,
        similar_threshold=0.8,
    )

    xlsx_path = output / "SOQL調査.xlsx"
    assert xlsx_path.exists()
    sheets = _read_xlsx(xlsx_path)

    # 一覧シート
    overview = sheets["一覧"]
    by_key = {row["管理番号"]: row for row in overview}
    assert by_key["1001"]["2000件超"] == "超えている"
    assert by_key["1002"]["2000件超"] == "超えていない"
    assert by_key["1002"]["行数"] == "10"
    # Contact.Email が 出力列にあるので 個人情報 は あり
    assert by_key["1001"]["個人情報"] == "あり"
    assert by_key["1002"]["個人情報"] == "あり"
    # 該当列に 「 メール 」 が出る
    assert "メール" in by_key["1001"]["個人情報の該当列"]
    # 似たレポートの列に 「 グループN: ... が基準 」 の形で出る
    assert "基準" in by_key["1001"]["似たレポート"]
    assert "基準" in by_key["1002"]["似たレポート"]

    # 列の対応シート: Contact.Email も出る
    columns_sheet = sheets["列の対応"]
    assert any(r["列キー"] == "Contact.Email" for r in columns_sheet)

    # 列名の対応.csv も作られている
    assert (output / "列名の対応.csv").exists()


def test_e2e_single_record_no_similar_text(tmp_path: Path) -> None:
    """1 件だけのときは 似たレポート列が空文字。 シート構成は 3 枚のまま。"""
    from src import pii
    from src.fetch import RELATED_DEPTH, RELATED_MAX, ROW_LIMIT_ENABLED, run_fetch
    from src.master import MasterEntry
    from src.tables import run_tables

    describe_csv_path = tmp_path / "describe.csv"
    pii_keywords = pii.DEFAULT_KEYWORDS
    pii_person_objects = pii.DEFAULT_PERSON_OBJECTS
    client = _FakeClient(
        describe=_describe("Opportunity", ["Opp.Name"]),
        object_describes=_object_describes(),
    )
    entry = MasterEntry(
        key="1001",
        summary="テスト",
        url=f"{DOMAIN}/00O000000000001/view",
        enabled=True,
    )
    outcome = run_fetch(
        [entry],
        related_depth=RELATED_DEPTH,
        related_max=RELATED_MAX,
        row_limit=ROW_LIMIT_ENABLED,
        pii_keywords=pii_keywords,
        pii_person_objects=pii_person_objects,
        describe_csv_path=describe_csv_path,
        site_for=_site_for(client),
    )[0]
    output = tmp_path / "output"
    output.mkdir(parents=True, exist_ok=True)
    run_tables(
        output,
        [outcome.record],
        similar_threshold=0.8,
    )
    assert (output / "SOQL調査.xlsx").exists()
    sheets = _read_xlsx(output / "SOQL調査.xlsx")
    assert len(sheets["一覧"]) == 1
    # 1 件だけのときは 似たレポートの列が空
    assert sheets["一覧"][0]["似たレポート"] == ""


def test_e2e_columns_sheet_resolves_related_email(tmp_path: Path) -> None:
    """``列の対応`` シートで、 関連経由（ Contact__r.Email ） の列も解決済みで出る。"""
    from src import pii
    from src.fetch import RELATED_DEPTH, RELATED_MAX, ROW_LIMIT_ENABLED, run_fetch
    from src.master import MasterEntry
    from src.tables import run_tables

    describe_csv_path = tmp_path / "describe.csv"
    pii_keywords = pii.DEFAULT_KEYWORDS
    pii_person_objects = pii.DEFAULT_PERSON_OBJECTS
    client = _FakeClient(
        describe=_describe(
            "Opportunity", ["Contact.Email"], column_labels={"Contact.Email": "メール"}
        ),
        object_describes=_object_describes(),
    )
    entry = MasterEntry(
        key="1001",
        summary="テスト",
        url=f"{DOMAIN}/00O000000000001/view",
        enabled=True,
    )
    outcome = run_fetch(
        [entry],
        related_depth=RELATED_DEPTH,
        related_max=RELATED_MAX,
        row_limit=ROW_LIMIT_ENABLED,
        pii_keywords=pii_keywords,
        pii_person_objects=pii_person_objects,
        describe_csv_path=describe_csv_path,
        site_for=_site_for(client),
    )[0]
    output = tmp_path / "output"
    output.mkdir(parents=True, exist_ok=True)
    run_tables(
        output,
        [outcome.record],
        similar_threshold=0.8,
    )
    sheets = _read_xlsx(output / "SOQL調査.xlsx")
    columns = sheets["列の対応"]
    # Contact.Email が出る
    email_rows = [r for r in columns if r["列キー"] == "Contact.Email"]
    assert email_rows
    # 使う SOQL 列名に Contact__r.Email のような関連経由が入る
    soql = email_rows[0]["使うSOQL列名"]
    assert "Email" in soql
