"""通し（ e2e ） テスト: ``run_fetch`` → ``run_tables`` → 出力 CSV の検証。

脳内スクラッチパッドの再現シナリオ （ Task / Account / User の偽 describe
と、 列 SUBJECT / ACCOUNT.NAME / OWNER.NAME、 絞り込み 2 件、 組立式
"1 OR 2"、 並び順、 topRows を持つレポート） を ``run_fetch`` から
``run_tables`` まで通して動かし、 出力 CSV の内容を厳密に確かめる。

ネットワークには出ない。 ``site_for`` を偽のクライアントに差し替える。
"""

from __future__ import annotations

import csv
import dataclasses
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table as XlsxTable
from openpyxl.worksheet.table import TableStyleInfo

DOMAIN = "https://example.my.salesforce.com/lightning/r/Report"

# ── ヘルパー ─────────────────────────────────────────────────────────────


def _build_master(path: Path) -> None:
    """``load_settings`` が困らないよう ``MASTER_XLSX_PATH`` の指す xlsx を作る。"""
    workbook = Workbook()
    active = workbook.active
    if active is not None:
        workbook.remove(active)
    sheet = workbook.create_sheet("PY_管理表")
    sheet.append(["ID", "概要", "Salesforce URL", "有効"])
    table = XlsxTable(displayName="PY_T_ReportEntry", ref="A1:D2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(path)


def _write_ini(tmp_path: Path, master: Path) -> None:
    """テスト用 ``config.ini`` を ``tmp_path`` に書く。"""
    (tmp_path / "config.ini").write_text(
        f"[FILES]\nMASTER_XLSX_PATH = {master.name}\nOUTPUT_DIR = ./output\n",
        encoding="utf-8",
    )


def _f(
    name: str,
    label: str,
    type_: str = "string",
    ref: list[str] | None = None,
    rel: str | None = None,
) -> dict:
    """``_slim_object`` を通す前の describe 用 ``field`` dict を作る。"""
    return {
        "name": name,
        "label": label,
        "type": type_,
        "custom": False,
        "referenceTo": ref or [],
        "relationshipName": rel,
        "picklistValues": [],
    }


# ── 偽物のクライアント ─────────────────────────────────────────────────


class _ReportStub:
    def __init__(self, describe: dict) -> None:
        self._describe = describe

    def describe(self, report_id: str) -> dict:
        return self._describe

    def get(self, report_id: str, filters: object = None, allow_truncated: bool = False) -> object:
        return []


class _FakeClient:
    def __init__(
        self,
        describe: dict,
        object_describes: dict[str, dict] | None = None,
    ) -> None:
        self.report = _ReportStub(describe)
        self._object_describes = object_describes or {}

    def describe_object(self, name: str) -> dict:
        return self._object_describes.get(name, {"name": name, "fields": []})


class _FakeSite:
    def __init__(self, client: _FakeClient) -> None:
        self._client = client

    def __call__(self) -> _FakeSite:
        return self

    def __enter__(self) -> _FakeClient:
        return self._client

    def __exit__(self, *args: object) -> None:
        return None


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """CSV を ``(columns, rows)`` で読む。"""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        columns = next(reader)
        rows = [dict(zip(columns, line, strict=True)) for line in reader]
    return columns, rows


def _site_for(client: _FakeClient):
    """``site_for(url)`` 互換の関数。"""

    def resolver(url: str) -> _FakeSite:
        return _FakeSite(client)

    return resolver


# ── 共通シナリオ ───────────────────────────────────────────────────────


def _e2e_describe() -> dict:
    """脳内スクラッチパッドの再現シナリオのレポート describe を返す。"""
    return {
        "reportMetadata": {
            "id": "00O000000000001",
            "name": "テスト",
            "reportType": {"type": "Task"},
            "reportFormat": "TABULAR",
            "detailColumns": ["SUBJECT", "ACCOUNT.NAME", "OWNER.NAME"],
            "reportFilters": [
                {"column": "SUBJECT", "operator": "contains", "value": "A"},
                {"column": "ACCOUNT.NAME", "operator": "equals", "value": "B"},
            ],
            "reportBooleanFilter": "1 OR 2",
            "sortBy": [{"sortColumn": "SUBJECT", "sortOrder": "Asc"}],
            "topRows": {"rowLimit": 10, "direction": "Asc"},
        },
        "reportExtendedMetadata": {
            "detailColumnInfo": {
                "SUBJECT": {"label": "件名"},
                "ACCOUNT.NAME": {"label": "取引先名"},
                "OWNER.NAME": {"label": "氏名"},
            }
        },
    }


def _e2e_object_describes() -> dict[str, dict]:
    """``Task`` / ``Account`` / ``User`` の偽 describe を返す。"""
    return {
        "Task": {
            "name": "Task",
            "label": "活動",
            "custom": False,
            "fields": [
                _f("Subject", "件名"),
                _f("WhatId", "関連先", "reference", ["Account"], "What"),
                _f("OwnerId", "担当者", "reference", ["User"], "Owner"),
            ],
        },
        "Account": {
            "name": "Account",
            "label": "取引先",
            "custom": False,
            "fields": [
                _f("Name", "取引先名"),
                _f("Phone", "電話", "phone"),
                _f("OwnerId", "所有者", "reference", ["User"], "Owner"),
                _f("ParentId", "親", "reference", ["Account"], "Parent"),
            ],
        },
        "User": {
            "name": "User",
            "label": "ユーザ",
            "custom": False,
            "fields": [
                _f("Name", "氏名"),
                _f("Email", "メール", "email"),
                _f("ManagerId", "マネージャ", "reference", ["User"], "Manager"),
            ],
        },
    }


def _run_e2e(tmp_path: Path) -> Path:
    """再現シナリオを実行し、 出力ディレクトリを返す。"""
    from src.fetch import run_fetch
    from src.master import MasterEntry
    from src.settings import load_settings
    from src.tables import run_tables

    master = tmp_path / "master.xlsx"
    _build_master(master)
    _write_ini(tmp_path, master)

    settings = load_settings(project_root=tmp_path)
    settings = dataclasses.replace(settings, output_dir=tmp_path / "output")

    client = _FakeClient(_e2e_describe(), _e2e_object_describes())

    entry = MasterEntry(
        key="1001",
        summary="テスト",
        url=f"{DOMAIN}/00O000000000001/view",
        enabled=True,
    )
    outcomes = run_fetch(settings, [entry], site_for=_site_for(client))
    assert all(o.status == "ok" for o in outcomes), [(o.status, o.error) for o in outcomes]
    run_tables(settings.output_dir, [o.record for o in outcomes if o.record])
    return settings.output_dir


# ── テスト本体 ───────────────────────────────────────────────────────────


def test_e2e_condition_table_項目_is_resolved(tmp_path: Path) -> None:
    """条件表の ``絞り込み`` 行で ``項目`` が column_map から解決されて埋まる。

    - ``SUBJECT`` → ``Subject``
    - ``ACCOUNT.NAME`` → ``What.Name`` （ 関連経由 ）
    """
    output_dir = _run_e2e(tmp_path)

    _, rows = _read_csv(output_dir / "条件表.csv")
    by_col: dict[str, dict[str, str]] = {}
    for row in rows:
        if row["種別"] == "絞り込み":
            by_col[row["列キー"]] = row

    assert by_col["SUBJECT"]["項目"] == "Subject"
    assert by_col["ACCOUNT.NAME"]["項目"] == "What.Name"
    # 演算子と値もそのまま
    assert by_col["SUBJECT"]["演算子"] == "contains"
    assert by_col["SUBJECT"]["値"] == "A"
    assert by_col["ACCOUNT.NAME"]["演算子"] == "equals"
    assert by_col["ACCOUNT.NAME"]["値"] == "B"

    # 組立式もそのまま
    bool_row = next(row for row in rows if row["種別"] == "組立式")
    assert bool_row["値"] == "1 OR 2"


def test_e2e_aggregate_table_sort_by_and_top_rows(tmp_path: Path) -> None:
    """集計表の ``並び順`` は ``sortColumn`` を読んで ``名前`` / ``内容`` に分け、
    条件表の ``上位N`` は dict を ``"N件 direction"`` の形に整形する。
    """
    output_dir = _run_e2e(tmp_path)

    # 集計表: 並び順 / 形式
    _, agg_rows = _read_csv(output_dir / "集計表.csv")
    agg_by_kind: dict[str, list[dict[str, str]]] = {}
    for row in agg_rows:
        agg_by_kind.setdefault(row["種別"], []).append(row)

    sort_row = agg_by_kind["並び順"][0]
    assert sort_row["名前"] == "SUBJECT"
    assert sort_row["内容"] == "Asc"
    # 順は 1
    assert sort_row["順"] == "1"

    fmt_row = agg_by_kind["形式"][0]
    assert fmt_row["内容"] == "TABULAR"

    # 条件表: 上位N
    _, cond_rows = _read_csv(output_dir / "条件表.csv")
    cond_by_kind: dict[str, list[dict[str, str]]] = {}
    for row in cond_rows:
        cond_by_kind.setdefault(row["種別"], []).append(row)

    top_row = cond_by_kind["上位N"][0]
    assert top_row["値"] == "10件 Asc"
    # Python repr 残っていない
    assert "{" not in top_row["値"]
    assert "'" not in top_row["値"]


def test_e2e_field_table_段_is_filled(tmp_path: Path) -> None:
    """項目表の ``段`` は ``relations`` と ``main_object`` から埋まる。

    - ``Task`` → 0 （ 主オブジェクト ）
    - ``Account`` → 1 （ Task.WhatId の参照先 ）
    - ``User`` → 1 （ Task.OwnerId の参照先 ）
    """
    output_dir = _run_e2e(tmp_path)

    _, rows = _read_csv(output_dir / "項目表.csv")
    by_obj: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_obj.setdefault(row["オブジェクト"], []).append(row)

    assert all(row["段"] == "0" for row in by_obj["Task"])
    assert all(row["段"] == "1" for row in by_obj["Account"])
    assert all(row["段"] == "1" for row in by_obj["User"])


def test_e2e_related_table(tmp_path: Path) -> None:
    """関連表は ``Task → Account / User`` と自己参照・ 子からの参照を
    全部載せる。
    """
    output_dir = _run_e2e(tmp_path)

    _, rows = _read_csv(output_dir / "関連表.csv")
    by_pair: dict[tuple[str, str, str], dict[str, str]] = {}
    for row in rows:
        key = (
            row["親オブジェクト"],
            row["参照項目API名"],
            row["子オブジェクト"],
        )
        by_pair[key] = row

    # Task 起点の 2 本
    assert ("Task", "WhatId", "Account") in by_pair
    assert ("Task", "OwnerId", "User") in by_pair
    # Account 起点の 2 本
    assert ("Account", "OwnerId", "User") in by_pair
    assert ("Account", "ParentId", "Account") in by_pair  # 自己参照
    # User 起点の 1 本
    assert ("User", "ManagerId", "User") in by_pair  # 自己参照

    # 段は全て 1 （ 関連は 1 段目 ）
    for row in rows:
        assert row["段"] == "1"


def test_e2e_column_names_csv(tmp_path: Path) -> None:
    """``列名の対応.csv`` は解決できた列だけが出る。"""
    output_dir = _run_e2e(tmp_path)

    _, rows = _read_csv(output_dir / "列名の対応.csv")
    soql_names = {row["SOQL列名"] for row in rows}
    assert soql_names == {"Subject", "What.Name", "Owner.Name"}
    # (不明) は出ない
    for row in rows:
        assert row["SOQL列名"] != "(不明)"
    # 管理番号と概要が埋まる
    for row in rows:
        assert row["管理番号"] == "1001"
        assert row["概要"] == "テスト"


def test_e2e_column_names_txt_full_text(tmp_path: Path) -> None:
    """``列名の対応_1001.txt`` の全文を厳密に検証する。

    期待する 4 スペースインデント、 ヘッダ、 dict 本文、 末尾の ``}`` まで
    そのまま一致するかを見る。
    """
    output_dir = _run_e2e(tmp_path)

    text = (output_dir / "列名の対応_1001.txt").read_text(encoding="utf-8")
    expected = (
        "# 管理番号: 1001\n"
        "# 概要: テスト\n"
        "# レポートID: 00O000000000001\n"
        "COLUMN_NAMES = {\n"
        '    "Subject": "件名",\n'
        '    "What.Name": "取引先名",\n'
        '    "Owner.Name": "氏名",\n'
        "}\n"
    )
    assert text == expected


# ── 集計表 / 条件表の行が「読める文字列」 になる網羅テスト ──────────────


def test_e2e_full_aggregate_table_all_kinds_are_human_readable(tmp_path: Path) -> None:
    """``groupingsDown`` / ``aggregates`` / ``buckets`` / ``customSummaryFormula``
    を持つレポートでも、 集計表の ``内容`` / ``名前`` が Python repr ではなく
    人が読める文字列になる。
    """
    from src.fetch import run_fetch
    from src.master import MasterEntry
    from src.settings import load_settings
    from src.tables import run_tables

    master = tmp_path / "master.xlsx"
    _build_master(master)
    _write_ini(tmp_path, master)
    settings = dataclasses.replace(
        load_settings(project_root=tmp_path), output_dir=tmp_path / "output"
    )

    describe = {
        "reportMetadata": {
            "reportType": {"type": "Task"},
            "reportFormat": "MATRIX",
            "detailColumns": ["SUBJECT"],
            "groupingsDown": [
                {
                    "name": "Opp.CloseDate",
                    "sortOrder": "Asc",
                    "dateGranularity": "FISCAL_QUARTER",
                },
            ],
            "groupingsAcross": [
                {"name": "Opp.StageName", "sortOrder": "Desc", "dateGranularity": "DAY"},
            ],
            "aggregates": ["s!Opp.Amount", "n!Opp.RollUp"],
            "sortBy": [{"sortColumn": "SUBJECT", "sortOrder": "Asc"}],
            "buckets": [
                {
                    "developerName": "AmountBucket",
                    "label": "金額帯",
                    "bucketType": "number",
                },
            ],
            "customSummaryFormula": [
                {"formula": "RowCount", "label": "件数", "dataType": "number"},
            ],
        },
        "reportExtendedMetadata": {
            "detailColumnInfo": {"SUBJECT": {"label": "件名"}},
        },
    }
    client = _FakeClient(describe, _e2e_object_describes())
    entry = MasterEntry(
        key="1001",
        summary="テスト",
        url=f"{DOMAIN}/00O000000000001/view",
        enabled=True,
    )
    outcomes = run_fetch(settings, [entry], site_for=_site_for(client))
    run_tables(settings.output_dir, [o.record for o in outcomes if o.record])

    _, rows = _read_csv(settings.output_dir / "集計表.csv")
    by_kind: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_kind.setdefault(row["種別"], []).append(row)

    # groupingsDown
    g_row = by_kind["グルーピング(行)"][0]
    assert g_row["名前"] == "Opp.CloseDate"
    assert g_row["内容"] == "Asc"
    # groupingsAcross
    ga_row = by_kind["グルーピング(列)"][0]
    assert ga_row["名前"] == "Opp.StageName"
    assert ga_row["内容"] == "Desc"
    # aggregates
    agg_names = [row["名前"] for row in by_kind["集計"]]
    assert agg_names == ["s!Opp.Amount", "n!Opp.RollUp"]
    # buckets
    b_row = by_kind["バケット"][0]
    assert b_row["名前"] == "AmountBucket"
    assert "金額帯" in b_row["内容"]
    assert "number" in b_row["内容"]
    # customSummaryFormula
    csf_row = by_kind["独自計算式"][0]
    assert csf_row["名前"] == "RowCount"
    assert csf_row["内容"] == "件数"

    # どの ``内容`` にも Python repr の残骸 （ ``{`` / ``'}`` ） が無いこと
    for row in rows:
        assert "{" not in row["名前"]
        assert "{" not in row["内容"] or row["種別"] in ("バケット",), f"Python repr っぽい: {row}"


def test_e2e_full_condition_table_all_kinds_are_human_readable(tmp_path: Path) -> None:
    """``standardDateFilter`` / ``standardFilters`` / ``crossFilters`` /
    ``topRows`` を持つレポートでも、 条件表の ``値`` が Python repr では
    なく人が読める文字列になる。
    """
    from src.fetch import run_fetch
    from src.master import MasterEntry
    from src.settings import load_settings
    from src.tables import run_tables

    master = tmp_path / "master.xlsx"
    _build_master(master)
    _write_ini(tmp_path, master)
    settings = dataclasses.replace(
        load_settings(project_root=tmp_path), output_dir=tmp_path / "output"
    )

    describe = {
        "reportMetadata": {
            "reportType": {"type": "Task"},
            "reportFormat": "TABULAR",
            "detailColumns": ["SUBJECT"],
            "reportFilters": [
                {"column": "SUBJECT", "operator": "contains", "value": "A"},
            ],
            "reportBooleanFilter": "1",
            "standardDateFilter": {
                "column": "Opp.CloseDate",
                "durationValue": "LAST_N_DAYS:30",
                "startDate": "2024-01-01",
                "endDate": "2024-01-31",
            },
            "standardFilters": [
                {"name": "Account.Industry", "value": "Banking"},
            ],
            "crossFilters": [
                {
                    "primaryEntityField": "Opp.AccountId",
                    "operator": "WITH",
                    "relatedEntity": "OpportunityContactRole",
                    "relatedEntityJoinField": "OpportunityId",
                    "criteria": [{"column": "Role", "operator": "equals", "value": "X"}],
                }
            ],
            "topRows": {"rowLimit": 10, "direction": "Asc"},
        },
        "reportExtendedMetadata": {
            "detailColumnInfo": {"SUBJECT": {"label": "件名"}},
        },
    }
    client = _FakeClient(describe, _e2e_object_describes())
    entry = MasterEntry(
        key="1001",
        summary="テスト",
        url=f"{DOMAIN}/00O000000000001/view",
        enabled=True,
    )
    outcomes = run_fetch(settings, [entry], site_for=_site_for(client))
    run_tables(settings.output_dir, [o.record for o in outcomes if o.record])

    _, rows = _read_csv(settings.output_dir / "条件表.csv")
    by_kind: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_kind.setdefault(row["種別"], []).append(row)

    # 絞り込みの ``項目`` は column_map から
    assert by_kind["絞り込み"][0]["項目"] == "Subject"
    # 組立式
    assert by_kind["組立式"][0]["値"] == "1"
    # 日付
    date_row = by_kind["日付"][0]
    assert date_row["列キー"] == "Opp.CloseDate"
    assert date_row["値"] == "2024-01-01〜2024-01-31"
    assert date_row["備考"] == "LAST_N_DAYS:30"
    # 標準フィルタ
    std_row = by_kind["標準フィルタ"][0]
    assert "{'name'" not in std_row["値"]
    assert '"name": "Account.Industry"' in std_row["値"]
    # クロス
    cross_row = by_kind["クロス"][0]
    assert "WITH" in cross_row["値"]
    assert "OpportunityId" in cross_row["値"]
    assert "{'column'" not in cross_row["値"]
    assert '"Role"' in cross_row["値"]
    # 上位N
    top_row = by_kind["上位N"][0]
    assert top_row["値"] == "10件 Asc"
    assert "{" not in top_row["値"]
