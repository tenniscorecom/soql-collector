"""``trial/`` の単品サンプルを fake client で検証する。

実 Salesforce に繋がないよう、 ``site_for`` をテストで差し替えられる形にし、
``tmp_path`` 配下の ``config.ini`` + 空の管理表で ``load_settings`` を動かす。

``trial/`` は **名前空間パッケージ** （ ``__init__.py`` を **持たない** ）。
``python -m trial.fetch_one_report`` でそのまま動くことを前提に、
``tests/`` 側での ``sys.path`` 書き換えは **行わない** 。
"""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table as XlsxTable
from openpyxl.worksheet.table import TableStyleInfo

DOMAIN = "https://example--sandbox.sandbox.my.salesforce.com/lightning/r/Report"


# ── ヘルパー ─────────────────────────────────────────────────────────────


def _build_master(path: Path) -> None:
    """``load_settings`` が困らないよう ``MASTER_XLSX_PATH`` の指す xlsx を作る。

    中身は空の ``PY_管理表`` 1 枚でよい（ 例題は管理表を読まない ）。
    """
    workbook = Workbook()
    active = workbook.active
    if active is not None:
        workbook.remove(active)
    sheet = workbook.create_sheet("PY_管理表")
    sheet.append(["ID", "概要", "Salesforce URL", "有効"])
    table = XlsxTable(displayName="PY_T_ReportEntry", ref="A1:D2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(path)


def _write_ini(tmp_path: Path, master: Path) -> None:
    """テスト用 ``config.ini`` を ``tmp_path`` に書く。"""
    (tmp_path / "config.ini").write_text(
        f"[FILES]\nMASTER_XLSX_PATH = {master.name}\nOUTPUT_DIR = ./output\n",
        encoding="utf-8",
    )


def _patch_load_settings(
    monkeypatch: pytest.MonkeyPatch, module: object, project_root: Path
) -> None:
    """``module.load_settings`` を ``project_root`` からの読み込みに差し替える。"""
    from src import settings as settings_module

    real_load = settings_module.load_settings

    def patched() -> object:
        return real_load(project_root=project_root)

    monkeypatch.setattr(module, "load_settings", patched)


# ── 偽物のクライアント ─────────────────────────────────────────────────


class _ReportStub:
    def __init__(self, describe: dict) -> None:
        self._describe = describe

    def describe(self, report_id: str) -> dict:
        return self._describe


class _FakeClient:
    def __init__(
        self,
        describe: dict,
        object_describes: dict[str, dict] | None = None,
    ) -> None:
        self.report = _ReportStub(describe)
        self._object_describes = object_describes or {}
        self.describe_object_calls: list[str] = []

    def describe_object(self, name: str) -> dict:
        self.describe_object_calls.append(name)
        return self._object_describes.get(name, {"name": name, "fields": []})

    def __enter__(self) -> _FakeClient:
        return self._client

    def __exit__(self, *args: object) -> None:
        return None


class _FakeSite:
    def __init__(self, client: _FakeClient) -> None:
        self._client = client

    def __call__(self) -> _FakeSite:
        return self

    def __enter__(self) -> _FakeClient:
        return self._client

    def __exit__(self, *args: object) -> None:
        return None


def _make_metadata(report_type: str = "Opportunity") -> dict:
    return {
        "reportMetadata": {
            "reportType": {"type": report_type},
            "reportFormat": "TABULAR",
            "detailColumns": ["Opp.Name"],
        },
        "reportExtendedMetadata": {
            "detailColumnInfo": {"Opp.Name": {"label": "名前"}},
        },
    }


def _site_for(client: _FakeClient) -> Callable[[str], _FakeSite]:
    """``site_for(url)`` 互換の関数。"""

    def resolver(url: str) -> _FakeSite:
        return _FakeSite(client)

    return resolver


# ── fetch_one_report ────────────────────────────────────────────────────


def test_fetch_one_report_no_argv_exits_2() -> None:
    """引数なしで終了コード 2 (argparse が使い方を出して ``SystemExit``)。"""
    from trial import fetch_one_report

    with pytest.raises(SystemExit) as exc:
        fetch_one_report.main([])
    assert exc.value.code == 2


def test_fetch_one_report_writes_csv(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """URL と管理番号を渡すと ``output/single/対応表_{管理番号}.csv`` と
    ``output/single/列名の対応_{管理番号}.txt`` が書かれる。 JSON は書かれない。"""
    from trial import fetch_one_report

    master = tmp_path / "master.xlsx"
    _build_master(master)
    _write_ini(tmp_path, master)
    _patch_load_settings(monkeypatch, fetch_one_report, tmp_path)

    client = _FakeClient(
        describe=_make_metadata(),
        object_describes={"Opportunity": {"name": "Opportunity", "fields": []}},
    )

    code = fetch_one_report.main(
        [f"{DOMAIN}/00O5g00000AAAAA/view", "1001"],
        site_for=_site_for(client),
    )

    assert code == 0
    assert (tmp_path / "output" / "single" / "対応表_1001.csv").exists()
    assert (tmp_path / "output" / "single" / "列名の対応_1001.txt").exists()
    # JSON は書かれていない
    assert not (tmp_path / "output" / "single" / "1001.json").exists()
    # 本体 output には触らない
    assert not (tmp_path / "output" / "1001.csv").exists()


def test_fetch_one_report_default_key_is_sample(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """管理番号を省略すると既定 ``sample`` で書かれる。"""
    from trial import fetch_one_report

    master = tmp_path / "master.xlsx"
    _build_master(master)
    _write_ini(tmp_path, master)
    _patch_load_settings(monkeypatch, fetch_one_report, tmp_path)

    client = _FakeClient(
        describe=_make_metadata(),
        object_describes={"Opportunity": {"name": "Opportunity", "fields": []}},
    )

    code = fetch_one_report.main(
        [f"{DOMAIN}/00O5g00000AAAAA/view"],
        site_for=_site_for(client),
    )

    assert code == 0
    assert (tmp_path / "output" / "single" / "対応表_sample.csv").exists()


def test_fetch_one_report_failure_returns_1(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """接続側 (``site_for``) が ``ComkenError`` を投げたら終了コード 1、 ファイルは作らない。

    ``run_fetch`` が ``ComkenError`` を握って ``FetchOutcome(status="failed")``
    に縮退する仕様に合わせて、 例外型を ``ComkenError`` にする。
    """
    from comken.exceptions import ComkenError

    from trial import fetch_one_report

    master = tmp_path / "master.xlsx"
    _build_master(master)
    _write_ini(tmp_path, master)
    _patch_load_settings(monkeypatch, fetch_one_report, tmp_path)

    def failing_site_for(url: str) -> _FakeSite:
        raise ComkenError("未登録組織")

    code = fetch_one_report.main(
        [f"{DOMAIN}/00O5g00000AAAAA/view", "1001"],
        site_for=failing_site_for,
    )

    assert code == 1
    assert not (tmp_path / "output" / "single").exists()
