"""レポート管理表（Excel）の読み取りテスト。

comken の ``Excel`` で読むので、実ファイルと同じ配置（``PY_管理表`` シート +
テーブル ``PY_T_ReportEntry``）を openpyxl で組み立てる。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.master import MasterEntry, _is_enabled, filter_enabled, read_master

DOMAIN = "https://example--sandbox.sandbox.my.salesforce.com/lightning/r/Report"


def _build_master(path: Path, rows: list[dict]) -> None:
    """テスト用の管理表 xlsx を作る（``PY_管理表`` シート + テーブル）。"""
    workbook = Workbook()
    # デフォルトで作られるアクティブシートを消す
    active = workbook.active
    if active is not None:
        workbook.remove(active)
    sheet = workbook.create_sheet("PY_管理表")
    headers = [
        "ID",
        "グループ名",
        "担当者",
        "概要",
        "Salesforce URL",
        "保存先",
        "有効",
        "0件あり",
        "備考",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(col, "") for col in headers])
    table_ref = f"A1:I{len(rows) + 1}"
    table = Table(displayName="PY_T_ReportEntry", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(path)


def test_read_master_basic(tmp_path: Path) -> None:
    path = tmp_path / "master.xlsx"
    _build_master(
        path,
        [
            {
                "ID": "1001",
                "概要": "顧客一覧",
                "Salesforce URL": f"{DOMAIN}/00O5g00000ABCDE/view",
                "有効": "○",
            },
            {
                "ID": "1002",
                "概要": "売上実績",
                "Salesforce URL": f"{DOMAIN}/00O5g00000FGHIJ/view",
                "有効": "×",
            },
        ],
    )

    entries = read_master(path)

    assert entries == [
        MasterEntry(
            key="1001", summary="顧客一覧", url=f"{DOMAIN}/00O5g00000ABCDE/view", enabled=True
        ),
        MasterEntry(
            key="1002", summary="売上実績", url=f"{DOMAIN}/00O5g00000FGHIJ/view", enabled=False
        ),
    ]


def test_read_master_blank_rows_skipped(tmp_path: Path) -> None:
    path = tmp_path / "master.xlsx"
    _build_master(
        path,
        [
            {"ID": "1001", "Salesforce URL": f"{DOMAIN}/00O5g00000ABCDE/view", "有効": "○"},
            {"ID": "", "Salesforce URL": "", "有効": ""},
            {"ID": "1002", "Salesforce URL": f"{DOMAIN}/00O5g00000FGHIJ/view", "有効": "×"},
        ],
    )

    entries = read_master(path)

    assert [entry.key for entry in entries] == ["1001", "1002"]


def test_read_master_enabled_variants(tmp_path: Path) -> None:
    path = tmp_path / "master.xlsx"
    _build_master(
        path,
        [
            {"ID": "1", "Salesforce URL": f"{DOMAIN}/00O5g00000AAAAA/view", "有効": "○"},
            {"ID": "2", "Salesforce URL": f"{DOMAIN}/00O5g00000BBBBB/view", "有効": "×"},
            {"ID": "3", "Salesforce URL": f"{DOMAIN}/00O5g00000CCCCC/view", "有効": "有効"},
            {"ID": "4", "Salesforce URL": f"{DOMAIN}/00O5g00000DDDDD/view", "有効": "はい"},
            {"ID": "5", "Salesforce URL": f"{DOMAIN}/00O5g00000EEEEE/view", "有効": ""},
        ],
    )

    entries = read_master(path)

    assert [entry.enabled for entry in entries] == [True, False, True, True, False]


def test_filter_enabled_drops_disabled(tmp_path: Path) -> None:
    path = tmp_path / "master.xlsx"
    _build_master(
        path,
        [
            {"ID": "1", "Salesforce URL": f"{DOMAIN}/00O5g00000AAAAA/view", "有効": "○"},
            {"ID": "2", "Salesforce URL": f"{DOMAIN}/00O5g00000BBBBB/view", "有効": "×"},
        ],
    )

    enabled = filter_enabled(read_master(path))
    assert [entry.key for entry in enabled] == ["1"]


def test_read_master_bad_url_kept_as_text(tmp_path: Path) -> None:
    """URL が壊れた行は、Excel 読み取り段階ではエラーにしない（``fetch`` 側で個別失敗）。"""
    path = tmp_path / "master.xlsx"
    _build_master(
        path,
        [
            {"ID": "1", "Salesforce URL": "壊れたURL", "有効": "○"},
            {"ID": "2", "Salesforce URL": f"{DOMAIN}/00O5g00000AAAAA/view", "有効": "○"},
        ],
    )

    entries = read_master(path)

    assert entries[0].url == "壊れたURL"
    assert entries[1].url.endswith("/00O5g00000AAAAA/view")


def test_is_enabled_aligned_with_downloader() -> None:
    """``_is_enabled`` は英語の ``"true"`` と ``Salesforceレポートダウンローダー`` の
    ``_to_bool`` と同じ判定を返す。空白つき ``" ○ "`` のような表記も吸収する。
    # なぜ: 「有効」判定の仕様はダウンローダー側と1つしか持たない
    # べきで、ここの判定が違うと fetch --all の対象が食い違うため
    """
    # 真
    assert _is_enabled("○") is True
    assert _is_enabled(" ○ ") is True  #  strip して比較
    assert _is_enabled("有効") is True
    assert _is_enabled("true") is True
    assert _is_enabled("TRUE") is True
    assert _is_enabled("True") is True
    assert _is_enabled("o") is True
    assert _is_enabled("yes") is True
    assert _is_enabled("1") is True
    assert _is_enabled("on") is True
    assert _is_enabled("はい") is True

    # 偽
    assert _is_enabled("×") is False
    assert _is_enabled("") is False
    assert _is_enabled("   ") is False  # strip 後は空
    assert _is_enabled("無効") is False
    assert _is_enabled("false") is False
    assert _is_enabled("no") is False
    assert _is_enabled("0") is False


def test_read_master_handles_true_and_whitespace(tmp_path: Path) -> None:
    """``"true"`` / ``" ○ "`` （空白つき） も ``enabled=True`` で読む。
    ``test_is_enabled_aligned_with_downloader`` の挙動が ``read_master`` 全体でも
    保たれていることを確かめる。
    """
    path = tmp_path / "master.xlsx"
    _build_master(
        path,
        [
            {"ID": "1", "Salesforce URL": f"{DOMAIN}/00O5g00000AAAAA/view", "有効": "true"},
            {"ID": "2", "Salesforce URL": f"{DOMAIN}/00O5g00000BBBBB/view", "有効": " TRUE "},
            {"ID": "3", "Salesforce URL": f"{DOMAIN}/00O5g00000CCCCC/view", "有効": " ○ "},
        ],
    )

    flags = [entry.enabled for entry in read_master(path)]
    assert flags == [True, True, True]
