"""``SOQL調査.xlsx`` の書き出し（テーブル・列幅・状態の色）。"""

from __future__ import annotations

import logging
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from src import xlsx_style
from src.tables.constants import (
    COLUMNS_SHEET_COLUMNS,
    SHEET_COLUMNS,
    SHEET_OVERVIEW,
    SHEET_RELATION_DIAGRAM,
    SURVEY_COLUMNS,
    TABLE_COLUMNS,
    TABLE_DIAGRAM,
    TABLE_OVERVIEW,
)
from src.tables.diagram_rows import relation_diagram_headers

logger = logging.getLogger(__name__)


# 一覧 / 列の対応 シートの 1 行目に出す案内文
_ANNOUNCEMENT_TEXT = (
    "直すときは input/列の直し.xlsx の『直したSOQL列名』に書いて、もう一度実行してください。"
)


def write_xlsx(
    path: Path,
    *,
    survey: list[dict[str, str]],
    columns: list[dict[str, str]],
    relation_diagram: list[dict[str, str]],
) -> None:
    """``SOQL調査.xlsx`` を書く。 シート 3 枚、 すべて Excel のテーブル。"""
    if path.exists():
        path.unlink()
    path.parent.mkdir(parents=True, exist_ok=True)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    # デフォルトの Sheet を消す
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)

    # 一覧シート
    overview = workbook.create_sheet(SHEET_OVERVIEW)
    overview["A1"] = _ANNOUNCEMENT_TEXT
    overview["A1"].font = Font(bold=True)
    overview["A1"].alignment = Alignment(wrap_text=True)
    _write_table(
        overview,
        headers=SURVEY_COLUMNS,
        rows=survey,
        start_row=3,
        start_col=1,
        table_name=TABLE_OVERVIEW,
    )
    _adjust_column_widths(overview, SURVEY_COLUMNS, survey, start_row=3)
    overview.freeze_panes = "A3"

    # 列の対応シート
    cols_sheet = workbook.create_sheet(SHEET_COLUMNS)
    cols_sheet["A1"] = _ANNOUNCEMENT_TEXT
    cols_sheet["A1"].font = Font(bold=True)
    cols_sheet["A1"].alignment = Alignment(wrap_text=True)
    _write_table(
        cols_sheet,
        headers=COLUMNS_SHEET_COLUMNS,
        rows=columns,
        start_row=3,
        start_col=1,
        table_name=TABLE_COLUMNS,
    )
    _apply_status_fills(cols_sheet, COLUMNS_SHEET_COLUMNS, columns, start_row=3)
    _adjust_column_widths(cols_sheet, COLUMNS_SHEET_COLUMNS, columns, start_row=3)
    cols_sheet.freeze_panes = "A3"

    # 関係図シート
    diagram_sheet = workbook.create_sheet(SHEET_RELATION_DIAGRAM)
    diagram_headers = relation_diagram_headers(relation_diagram)
    _write_table(
        diagram_sheet,
        headers=diagram_headers,
        rows=relation_diagram,
        start_row=1,
        start_col=1,
        table_name=TABLE_DIAGRAM,
    )
    _adjust_column_widths(diagram_sheet, diagram_headers, relation_diagram, start_row=1)
    diagram_sheet.freeze_panes = "A2"

    workbook.save(path)


def _write_table(
    worksheet: Any,
    *,
    headers: Sequence[str],
    rows: list[dict[str, str]],
    start_row: int,
    start_col: int,
    table_name: str,
) -> None:
    """openpyxl で 「見出し行 + データ行 + Excel テーブル」 を一度に書く。"""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as _gcl
    from openpyxl.worksheet.table import Table as XlsxTable
    from openpyxl.worksheet.table import TableStyleInfo

    # 見出し
    for col_offset, header in enumerate(headers):
        cell = worksheet.cell(row=start_row, column=start_col + col_offset, value=header)
        cell.font = Font(bold=True)
    # データ
    for row_offset, row in enumerate(rows):
        for col_offset, header in enumerate(headers):
            worksheet.cell(
                row=start_row + 1 + row_offset,
                column=start_col + col_offset,
                value=row.get(header, ""),
            )
    end_row = start_row + len(rows)
    end_col = start_col + len(headers) - 1
    if end_row <= start_row:
        # データが無いときは Excel テーブルが作れないのでヘッダーだけにする
        return
    ref = f"{_gcl(start_col)}{start_row}:{_gcl(end_col)}{end_row}"
    table = XlsxTable(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _adjust_column_widths(
    worksheet: Any,
    headers: Sequence[str],
    rows: list[dict[str, str]],
    *,
    start_row: int,
) -> None:
    """``xlsx_style.adjust_column_widths`` への薄いラッパー。"""
    xlsx_style.adjust_column_widths(
        worksheet,
        headers,
        rows,
        start_row=start_row,
    )


def _apply_status_fills(
    worksheet: Any,
    headers: Sequence[str],
    rows: list[dict[str, str]],
    *,
    start_row: int,
) -> None:
    """状態列の値でセルを塗り分ける。 要対応=薄い赤、 要確認=薄い黄。"""
    from openpyxl.styles import PatternFill

    if "状態" not in headers:
        return
    # start_row は見出し行、 データは start_row+1 から
    fill_fix = PatternFill(
        start_color=xlsx_style.FILL_NEEDS_FIX,
        end_color=xlsx_style.FILL_NEEDS_FIX,
        fill_type="solid",
    )
    fill_check = PatternFill(
        start_color=xlsx_style.FILL_NEEDS_CHECK,
        end_color=xlsx_style.FILL_NEEDS_CHECK,
        fill_type="solid",
    )
    for row_offset, row in enumerate(rows):
        status = row.get("状態", "")
        if status == "要対応":
            fill = fill_fix
        elif status == "要確認":
            fill = fill_check
        else:
            continue
        for col_index in range(1, len(headers) + 1):
            cell = worksheet.cell(row=start_row + 1 + row_offset, column=col_index)
            cell.fill = fill
