"""Excel ブックと describe JSON の永続化。"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

SHEETS: dict[str, tuple[str, ...]] = {
    "Master": ("管理番号", "概要", "レポートID", "URL"),
    "Reports": (
        "管理番号",
        "概要",
        "レポートID",
        "URL",
        "組織",
        "状態",
        "SOQLドラフト",
        "備考",
        "最終取得日時",
        "JSONパス",
        "主オブジェクト",
        "reportFormat",
    ),
    "Filters": ("レポートID", "種別", "列", "演算子", "値", "表示順"),
    "Groupings": ("レポートID", "種別", "名前", "並び順"),
    "FieldMappings": (
        "サイト",
        "レポートタイプ",
        "列キー",
        "表示名",
        "フィールドAPI名",
        "型",
        "確認状態",
        "確認日",
        "確認者",
        "備考",
    ),
    "Notes": ("レポートID", "日時", "種別", "内容"),
}


class WorkbookStore:
    """6 シート構成の Excel ブックを行辞書として扱う。"""

    def __init__(self, path: Path) -> None:
        self.path = path
        self._ensure_book()

    def _ensure_book(self) -> None:
        if self.path.exists():
            return
        self.path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, headers in SHEETS.items():
            worksheet = workbook.create_sheet(name)
            worksheet.append(headers)
            worksheet.freeze_panes = "A2"
        workbook.save(self.path)

    def read_rows(self, sheet_name: str) -> list[dict[str, str]]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            headers = [str(cell.value or "") for cell in worksheet[1]]
            return [
                {header: "" if value is None else str(value) for header, value in zip(headers, row)}
                for row in worksheet.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)
            ]
        finally:
            workbook.close()

    def replace_rows(self, sheet_name: str, rows: Iterable[dict[str, object]]) -> None:
        workbook = load_workbook(self.path)
        try:
            worksheet = workbook[sheet_name]
            if worksheet.max_row > 1:
                worksheet.delete_rows(2, worksheet.max_row - 1)
            headers = SHEETS[sheet_name]
            for row in rows:
                worksheet.append([row.get(header, "") for header in headers])
            workbook.save(self.path)
        finally:
            workbook.close()

    def append_row(self, sheet_name: str, row: dict[str, object]) -> None:
        rows = self.read_rows(sheet_name)
        rows.append({key: str(value) for key, value in row.items()})
        self.replace_rows(sheet_name, rows)

    def upsert(self, sheet_name: str, key_columns: tuple[str, ...], row: dict[str, object]) -> None:
        rows = self.read_rows(sheet_name)
        key = tuple(str(row.get(column, "")) for column in key_columns)
        replacement = {column: str(row.get(column, "")) for column in SHEETS[sheet_name]}
        for index, current in enumerate(rows):
            if tuple(current[column] for column in key_columns) == key:
                rows[index] = replacement
                break
        else:
            rows.append(replacement)
        self.replace_rows(sheet_name, rows)

    def find_report(self, report_id: str) -> dict[str, str] | None:
        return next(
            (row for row in self.read_rows("Reports") if row["レポートID"] == report_id), None
        )

    def add_note(self, report_id: str, kind: str, content: str) -> None:
        self.append_row(
            "Notes", {"レポートID": report_id, "日時": now_text(), "種別": kind, "内容": content}
        )


def write_json(path: Path, metadata: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")


def now_text() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")
