"""CLI から呼ぶユースケース。"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from soql_collector.settings import Settings
from soql_collector.soql import Draft, compose_soql
from soql_collector.store import WorkbookStore, now_text, write_json

REPORT_ID_RE = re.compile(r"\b(00O[a-zA-Z0-9]{12,15})\b")
STATUSES = {"READY", "REVIEW", "BLOCKED", "INVALID", "ERROR", "PENDING"}
CSV_HEADERS = (
    "管理番号",
    "概要",
    "レポートID",
    "URL",
    "状態",
    "SOQLドラフト",
    "備考",
    "フィルタ詳細(生データ)",
    "集計・グルーピング詳細(生データ)",
)


def report_id_from_text(text: str) -> str:
    match = REPORT_ID_RE.search(text)
    return match.group(1) if match else text.strip()


def import_master(store: WorkbookStore, source: Path) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook["Master"] if "Master" in workbook.sheetnames else workbook.active
        headers = [str(cell.value or "") for cell in worksheet[1]]
        rows = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            source_row = dict(zip(headers, values))
            url = str(source_row.get("URL") or "")
            report_id = str(source_row.get("レポートID") or report_id_from_text(url))
            if report_id:
                rows.append(
                    {
                        "管理番号": source_row.get("管理番号", ""),
                        "概要": source_row.get("概要", ""),
                        "レポートID": report_id,
                        "URL": url,
                    }
                )
    finally:
        workbook.close()
    store.replace_rows("Master", rows)
    return len(rows)


def mappings_for(store: WorkbookStore, site: str, report_type: str) -> dict[str, tuple[str, str]]:
    return {
        row["列キー"]: (row["フィールドAPI名"], row["型"])
        for row in store.read_rows("FieldMappings")
        if row["サイト"] == site
        and row["レポートタイプ"] == report_type
        and row["確認状態"] == "確認済み"
    }


def build_saved_report(
    store: WorkbookStore, report_id: str, is_apply: bool = False
) -> Draft | None:
    report = store.find_report(report_id)
    if report is None or not report["JSONパス"] or not Path(report["JSONパス"]).exists():
        return None
    metadata = json.loads(Path(report["JSONパス"]).read_text(encoding="utf-8"))
    draft = compose_soql(metadata, mappings_for(store, report["組織"], report["主オブジェクト"]))
    if is_apply:
        report["SOQLドラフト"] = draft.soql
        report["備考"] = " / ".join(draft.notes)
        report["状態"] = "READY" if draft.is_complete else "BLOCKED"
        store.upsert("Reports", ("レポートID",), report)
    return draft


def collect_one(
    settings: Settings,
    store: WorkbookStore,
    url: str,
    org: str | None = None,
    is_dry_run: bool = False,
    client_factory: Callable | None = None,
) -> dict[str, str]:
    report_id = report_id_from_text(url)
    if client_factory is None:
        from comken.toolbox.salesforce.sites import site_for

        site_class = site_for(url)
        client_factory = site_class
        org = org or site_class.__name__
    if is_dry_run:
        return {"レポートID": report_id, "状態": "PENDING"}
    with client_factory() as client:
        metadata = client.report.describe(report_id)
        fields_table, object_error = client.report.describe_fields_with_object_status(metadata)
    report_metadata = metadata.get("reportMetadata", {})
    report_type = report_metadata.get("reportType", {})
    object_name = report_type.get("type", "") if isinstance(report_type, dict) else ""
    site_name = org or ""
    existing_mappings = {
        (row["サイト"], row["レポートタイプ"], row["列キー"]): row
        for row in store.read_rows("FieldMappings")
    }
    for field in fields_table:
        key = (site_name, object_name, field["列キー"])
        current = existing_mappings.get(key, {})
        is_confirmed = current.get("確認状態") == "確認済み"
        store.upsert(
            "FieldMappings",
            ("サイト", "レポートタイプ", "列キー"),
            {
                "サイト": site_name,
                "レポートタイプ": object_name,
                "列キー": field["列キー"],
                "表示名": field["表示名"],
                "フィールドAPI名": current.get("フィールドAPI名", "")
                if is_confirmed
                else field["フィールドAPI名"],
                "型": current.get("型", "") if is_confirmed else field["型"],
                "確認状態": "確認済み" if is_confirmed else "未確認",
                "確認日": current.get("確認日", ""),
                "確認者": current.get("確認者", ""),
                "備考": current.get("備考", "") if is_confirmed else field.get("備考", ""),
            },
        )
    draft = compose_soql(metadata, mappings_for(store, site_name, object_name))
    json_path = settings.json_dir / f"{report_id}.json"
    write_json(json_path, metadata)
    master = next((row for row in store.read_rows("Master") if row["レポートID"] == report_id), {})
    row = {
        "管理番号": master.get("管理番号", ""),
        "概要": master.get("概要", ""),
        "レポートID": report_id,
        "URL": url,
        "組織": site_name,
        "状態": "READY" if draft.is_complete else "BLOCKED",
        "SOQLドラフト": draft.soql if draft.is_complete else "",
        "備考": " / ".join(([object_error] if object_error else []) + list(draft.notes)),
        "最終取得日時": now_text(),
        "JSONパス": str(json_path.resolve()),
        "主オブジェクト": object_name,
        "reportFormat": report_metadata.get("reportFormat", ""),
    }
    store.upsert("Reports", ("レポートID",), row)
    _replace_details(store, report_id, report_metadata)
    store.add_note(report_id, "collect", "describe() を取得しました")
    return row


def _replace_details(store: WorkbookStore, report_id: str, metadata: dict) -> None:
    filters = [row for row in store.read_rows("Filters") if row["レポートID"] != report_id]
    for index, item in enumerate(metadata.get("reportFilters", []), 1):
        filters.append(
            {
                "レポートID": report_id,
                "種別": "filter",
                "列": item.get("column", ""),
                "演算子": item.get("operator", ""),
                "値": item.get("value", ""),
                "表示順": index,
            }
        )
    for index, item in enumerate(metadata.get("crossFilters", []), 1):
        filters.append(
            {
                "レポートID": report_id,
                "種別": "crossFilter",
                "列": item.get("relatedEntity", ""),
                "演算子": item.get("operation", ""),
                "値": str(item),
                "表示順": index,
            }
        )
    store.replace_rows("Filters", filters)
    groupings = [row for row in store.read_rows("Groupings") if row["レポートID"] != report_id]
    for kind in ("aggregates", "groupingsDown", "groupingsAcross"):
        for index, item in enumerate(metadata.get(kind, []), 1):
            name = item.get("name", item.get("field", "")) if isinstance(item, dict) else item
            groupings.append({"レポートID": report_id, "種別": kind, "名前": name, "並び順": index})
    store.replace_rows("Groupings", groupings)


def export_csv(store: WorkbookStore, path: Path) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = store.read_rows("Reports")
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in CSV_HEADERS})
    return len(rows)
